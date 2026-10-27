"""Selecção de 2009: quais dos 111 espaços fazem parte dos 50.

Duas listas: a «Lista Final» de 2009-04-02 (48 marcados P) e a lista do livro (50, mais recente).
A coluna numérica da «Lista Final» é só um contador por ordem alfabética de concelho, não uma
classificação: por isso não se extrai nem se publica ordem. Colunas dos grupos de avaliação
(pontuações individuais) não são lidas.
"""
import json
import unicodedata
from pathlib import Path

import openpyxl

from espacos_pipeline.fontes import PROJECTO as PROJ

RAIZ = Path(__file__).resolve().parents[1]
LISTA_FINAL = PROJ / "Espacos" / "Avaliacao" / "2009-04-02 - SeleccaoEV.xlsx"
LIVRO = PROJ / "Mais recente" / "50espacos - livro.xlsx"
SAIDA = RAIZ / "docs" / "fonte_campanha" / "extraido" / "selecao2009.json"


def _chave(nome: str) -> str:
    sem = unicodedata.normalize("NFKD", str(nome)).encode("ascii", "ignore").decode()
    return " ".join(sem.lower().split())


def ler_selecao(lista_final: Path, livro: Path) -> list:
    wb = openpyxl.load_workbook(lista_final, data_only=True, read_only=True)
    cab, *linhas = [r for r in wb["Lista Final"].iter_rows(values_only=True) if r[0] is not None]
    col = {str(nome).strip().upper(): i for i, nome in enumerate(cab) if nome}
    espacos = {}
    for r in linhas:
        nome = str(r[col["NOME"]]).strip()
        marca = str(r[col["LISTA FINAL"]] or "").strip().upper()
        espacos[_chave(nome)] = {
            "id": int(r[col["ID"]]), "nome": nome, "concelho": r[col["CONCELHO"]],
            "selecao": "2009-04" if marca == "P" else None, "area_livro_ha": None,
        }
    wb2 = openpyxl.load_workbook(livro, data_only=True, read_only=True)
    cab2, *linhas2 = [r for r in wb2["Espacos"].iter_rows(values_only=True) if r[1]]
    col2 = {str(nome): i for i, nome in enumerate(cab2) if nome}
    area_col = next(i for nome, i in col2.items() if nome.startswith(("Área", "�rea", "rea", "Area")) or "ha" in nome.lower())
    for r in linhas2:
        k = _chave(r[col2["Nome"]])
        if r[col2["Lista_final"]] != "Sim":
            continue
        e = espacos.setdefault(k, {"id": None, "nome": str(r[col2["Nome"]]).strip(), "concelho": r[col2["Concelho"]], "selecao": None, "area_livro_ha": None})
        if e["selecao"] is None:
            e["selecao"] = "posterior"
        e["area_livro_ha"] = r[area_col]
    return list(espacos.values())


if __name__ == "__main__":
    dados = ler_selecao(LISTA_FINAL, LIVRO)
    SAIDA.parent.mkdir(parents=True, exist_ok=True)
    SAIDA.write_text(json.dumps(dados, ensure_ascii=False, indent=1), encoding="utf-8")
    n = {s: sum(e["selecao"] == s for e in dados) for s in ("2009-04", "posterior")}
    print(f"{len(dados)} espaços; seleccionados {n}; total {sum(n.values())}")
